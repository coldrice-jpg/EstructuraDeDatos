import tkinter as tk
from tkinter import ttk, messagebox
import random, time
from docx import Document
from docx.shared import Inches

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import matplotlib.pyplot as plt


# GENERACIÓN DE ARREGLOS

def best_case(n): return list(range(1, n+1))
def worst_case(n): return list(range(n, 0, -1))
def average_case(n):
    arr = list(range(1, n+1))
    random.shuffle(arr)
    return arr

def user_array(txt):
    return list(map(int, txt.split(",")))


# ALGORITMOS (VISUALES)

def bubble_visual(arr):
    a = arr.copy()
    for i in range(len(a)):
        for j in range(len(a)-i-1):
            if a[j] > a[j+1]:
                a[j], a[j+1] = a[j+1], a[j]
                yield a

def insertion_visual(arr):
    a = arr.copy()
    for i in range(1, len(a)):
        key = a[i]
        j = i-1
        while j >= 0 and a[j] > key:
            a[j+1] = a[j]
            j -= 1
            yield a
        a[j+1] = key
        yield a

def quick_visual(arr):
    a = arr.copy()

    def qs(low, high):
        if low >= high:
            return
        pivot = a[(low + high) // 2]
        i, j = low, high

        while i <= j:
            while a[i] < pivot:
                i += 1
            while a[j] > pivot:
                j -= 1
            if i <= j:
                a[i], a[j] = a[j], a[i]
                yield a.copy()
                i += 1
                j -= 1

        yield from qs(low, j)
        yield from qs(i, high)

    yield from qs(0, len(a) - 1)


# ALGORITMOS (REALES)

def bubble_sort(arr):
    a = arr.copy(); ops = 0
    for i in range(len(a)):
        for j in range(len(a)-i-1):
            ops += 1
            if a[j] > a[j+1]:
                a[j], a[j+1] = a[j+1], a[j]
    return a, ops

def insertion_sort(arr):
    a = arr.copy(); ops = 0
    for i in range(1, len(a)):
        key = a[i]; j = i-1
        while j >= 0 and a[j] > key:
            ops += 1
            a[j+1] = a[j]
            j -= 1
        a[j+1] = key
    return a, ops

def quick_sort(arr):
    ops = 0
    def qs(a):
        nonlocal ops
        if len(a) <= 1: return a
        pivot = a[len(a)//2]
        left, mid, right = [], [], []
        for x in a:
            ops += 1
            if x < pivot: left.append(x)
            elif x == pivot: mid.append(x)
            else: right.append(x)
        return qs(left) + mid + qs(right)
    return qs(arr.copy()), ops


# MEDICIÓN

def benchmark(func, arr):
    start = time.perf_counter_ns()
    _, ops = func(arr)
    end = time.perf_counter_ns()
    return end-start, ops


# GUI PRINCIPAL CON SCROLL

root = tk.Tk()
root.title("Visualizador y Analizador de Algoritmos")
root.geometry("900x700")

canvas_root = tk.Canvas(root)
scroll = ttk.Scrollbar(root, orient="vertical", command=canvas_root.yview)
scroll.pack(side="right", fill="y")
canvas_root.pack(side="left", fill="both", expand=True)

frame = ttk.Frame(canvas_root)
canvas_root.create_window((0,0), window=frame, anchor="nw")
canvas_root.configure(yscrollcommand=scroll.set)


# CONTROLES

algorithms = {
    "Bubble Sort": (bubble_sort, bubble_visual),
    "Insertion Sort": (insertion_sort, insertion_visual),
    "Quick Sort": (quick_sort, quick_visual)
}

algo1 = tk.StringVar(value="Bubble Sort")
algo2 = tk.StringVar(value="Insertion Sort")
case = tk.StringVar(value="Promedio")
n_size = tk.IntVar(value=30)
user_input = tk.StringVar()

controls = ttk.LabelFrame(frame, text="Configuración")
controls.pack(fill="x", pady=10)

ttk.Label(controls, text="Algoritmo A").grid(row=0, column=0)
ttk.Combobox(controls, textvariable=algo1, values=list(algorithms.keys())).grid(row=0, column=1)

ttk.Label(controls, text="Algoritmo B").grid(row=0, column=2)
ttk.Combobox(controls, textvariable=algo2, values=list(algorithms.keys())).grid(row=0, column=3)

ttk.Label(controls, text="Caso").grid(row=1, column=0)
ttk.Combobox(controls, textvariable=case,
            values=["Mejor","Promedio","Peor","Usuario"]).grid(row=1, column=1)

ttk.Label(controls, text="Tamaño N").grid(row=1, column=2)
ttk.Entry(controls, textvariable=n_size, width=6).grid(row=1, column=3)

ttk.Label(controls, text="Arreglo Usuario").grid(row=2, column=0)
ttk.Entry(controls, textvariable=user_input, width=40).grid(row=2, column=1, columnspan=3)


# VISUALIZACIÓN

vis_frame = ttk.Frame(frame)
vis_frame.pack(pady=10)

canvasA = tk.Canvas(vis_frame, width=350, height=200, bg="white")
canvasA.grid(row=0, column=0, padx=10)
canvasB = tk.Canvas(vis_frame, width=350, height=200, bg="white")
canvasB.grid(row=0, column=1, padx=10)

def draw(canvas, arr):
    canvas.delete("all")
    w,h = 350,200
    bar = w/len(arr)
    m = max(arr)
    for i,v in enumerate(arr):
        canvas.create_rectangle(i*bar, h-(v/m)*h, (i+1)*bar, h, fill="#4CAF50")

def animate(gen, canvas):
    try:
        arr = next(gen)
        draw(canvas, arr)
        root.after(40, animate, gen, canvas)
    except StopIteration:
        pass


# GRÁFICAS EMBEBIDAS

fig1, ax1 = plt.subplots(figsize=(6,3))
fig2, ax2 = plt.subplots(figsize=(6,3))

canvas_plot1 = FigureCanvasTkAgg(fig1, frame)
canvas_plot1.get_tk_widget().pack()
canvas_plot2 = FigureCanvasTkAgg(fig2, frame)
canvas_plot2.get_tk_widget().pack()


# EJECUCIÓN

results = {}

def run_all():
    N = [50,100,200,400]
    t1,t2,o1,o2 = [],[],[],[]

    for n in N:
        if case.get()=="Mejor": arr = best_case(n)
        elif case.get()=="Peor": arr = worst_case(n)
        elif case.get()=="Usuario": arr = user_array(user_input.get())
        else: arr = average_case(n)

        r1 = benchmark(algorithms[algo1.get()][0], arr)
        r2 = benchmark(algorithms[algo2.get()][0], arr)
        t1.append(r1[0]); t2.append(r2[0])
        o1.append(r1[1]); o2.append(r2[1])

    ax1.clear()
    ax1.plot(N,t1,label=algo1.get())
    ax1.plot(N,t2,label=algo2.get())
    ax1.set_title("Tiempo vs N"); ax1.legend()
    fig1.savefig("tiempo.png")
    canvas_plot1.draw()

    ax2.clear()
    ax2.plot(N,o1,label=algo1.get())
    ax2.plot(N,o2,label=algo2.get())
    ax2.set_title("Operaciones vs N"); ax2.legend()
    fig2.savefig("ops.png")
    canvas_plot2.draw()

    base = average_case(n_size.get())
    animate(algorithms[algo1.get()][1](base), canvasA)
    animate(algorithms[algo2.get()][1](base), canvasB)

def generate_doc():
    doc = Document()
    doc.add_heading("Reporte de Análisis de Algoritmos",1)

    doc.add_paragraph(
        "Este reporte presenta un análisis comparativo de algoritmos de ordenamiento "
        "evaluando su eficiencia en tiempo de ejecución y número de operaciones."
    )

    doc.add_heading("Algoritmos Comparados",2)
    doc.add_paragraph(f"Algoritmo A: {algo1.get()}")
    doc.add_paragraph(f"Algoritmo B: {algo2.get()}")

    doc.add_heading("Análisis de Tiempo",2)
    doc.add_paragraph(
        "La siguiente gráfica muestra el tiempo de ejecución conforme aumenta el tamaño "
        "del arreglo. Un crecimiento menor indica mejor eficiencia."
    )
    doc.add_picture("tiempo.png", width=Inches(5))

    doc.add_heading("Análisis de Operaciones",2)
    doc.add_paragraph(
        "Esta gráfica representa el número de operaciones realizadas por cada algoritmo, "
        "lo que permite analizar su complejidad computacional."
    )
    doc.add_picture("ops.png", width=Inches(5))

    doc.add_heading("Conclusión",2)
    doc.add_paragraph(
        "Con base en los resultados, se puede identificar qué algoritmo es más eficiente "
        "dependiendo del tamaño del arreglo y el caso analizado."
    )

    doc.save("Reporte_Final.docx")
    messagebox.showinfo("Reporte","Documento generado correctamente")

ttk.Button(frame, text="Ejecutar y Visualizar", command=run_all).pack(pady=10)
ttk.Button(frame, text="Generar Documento", command=generate_doc).pack(pady=20)

frame.update_idletasks()
canvas_root.configure(scrollregion=canvas_root.bbox("all"))

root.mainloop()
